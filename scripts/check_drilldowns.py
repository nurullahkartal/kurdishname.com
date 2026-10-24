import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

files = [
    ('gsc_export/kurdishname.com-Coverage-Drilldown-2026-10-12 (1).xlsx', 'Redirects'),
    ('gsc_export/kurdishname.com-Coverage-Drilldown-2026-10-12 (2).xlsx', 'Duplicate no canonical'),
    ('gsc_export/kurdishname.com-Coverage-Drilldown-2026-10-12 (3).xlsx', 'Crawled not indexed'),
    ('gsc_export/kurdishname.com-Coverage-Drilldown-2026-10-12 (6).xlsx', 'Discovered not indexed')
]

for f, name in files:
    try:
        wb = openpyxl.load_workbook(f)
        sheet = wb['Tablo']
        print(f"\n=== {name} ({sheet.max_row - 1} rows) ===")
        for i in range(2, min(12, sheet.max_row + 1)):
            print(' ', sheet.cell(row=i, column=1).value)
    except Exception as e:
        print(f"Error reading {f}: {e}")
